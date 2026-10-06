"""Render text blocks into real file formats (PDF, DOCX, HTML, EML, PNG, MD).

Blocks can be marked non-visible (e.g. white 1pt text in a PDF, a vanished run in DOCX,
display:none in HTML, an HTML-only part in an email) so the extractor (CORE-2) can be tested on
"what the model sees vs. what a human sees". Used for fixtures and evaluation cases alike.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path
from email.message import EmailMessage
from html import escape


@dataclass
class Block:
    text: str
    visible: bool = True


def _blocks(content: str | list[Block]) -> list[Block]:
    return [Block(content)] if isinstance(content, str) else content


def to_html(content: str | list[Block], title: str = "Document") -> str:
    parts = []
    for b in _blocks(content):
        body = escape(b.text).replace("\n", "<br>")
        parts.append(f"<p>{body}</p>" if b.visible else f'<div style="display:none">{body}</div>')
    return f"<html><head><title>{escape(title)}</title></head><body><h1>{escape(title)}</h1>{''.join(parts)}</body></html>"


def to_markdown(content: str | list[Block], title: str = "Document") -> str:
    out = [f"# {title}", ""]
    for b in _blocks(content):
        out.append(b.text if b.visible else f"<!-- {b.text} -->")
        out.append("")
    return "\n".join(out)


_FONT_CANDIDATES = [
    "C:/Windows/Fonts/arial.ttf",                                  # Windows
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",             # Debian/Ubuntu
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",                      # Fedora/RHEL
    "/Library/Fonts/Arial Unicode.ttf",                            # macOS
]


def _pdf_font() -> tuple[str, str]:
    """A Unicode TrueType font when one is installed, so non-Latin and look-alike letters render as themselves
    (the built-in Helvetica has no Cyrillic glyphs: such text reached readers as boxes, D-031). Falls back to
    Helvetica when no font is found."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "IS-Unicode" in pdfmetrics.getRegisteredFontNames():
        return "IS-Unicode", "IS-Unicode"
    for path in _FONT_CANDIDATES:
        if Path(path).exists():
            try:
                pdfmetrics.registerFont(TTFont("IS-Unicode", path))
                return "IS-Unicode", "IS-Unicode"
            except Exception:
                continue
    return "Helvetica", "Helvetica-Bold"


def to_pdf(content: str | list[Block], title: str = "Document") -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    body_font, title_font = _pdf_font()

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setTitle(title)
    w, h = A4
    y = h - 60
    c.setFont(title_font, 14)
    c.drawString(50, y, title)
    y -= 28
    for b in _blocks(content):
        if b.visible:
            c.setFillColorRGB(0, 0, 0)
            c.setFont(body_font, 10)
            size = 10
        else:
            c.setFillColorRGB(1, 1, 1)
            c.setFont(body_font, 1)
            size = 1
        for line in _wrap(b.text, 95):
            if y < 50:
                c.showPage()
                y = h - 50
            c.drawString(50, y, line)
            y -= max(size + 4, 3)
        y -= 6 if b.visible else 0
    c.save()
    return buf.getvalue()


def to_docx(content: str | list[Block], title: str = "Document") -> bytes:
    import docx

    d = docx.Document()
    d.add_heading(title, level=1)
    for b in _blocks(content):
        run = d.add_paragraph().add_run(b.text)
        if not b.visible:
            run.font.hidden = True
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def to_eml(content: str | list[Block], subject: str, sender: str, to: str = "user@example-corp.test") -> bytes:
    blocks = _blocks(content)
    m = EmailMessage()
    m["From"], m["To"], m["Subject"] = sender, to, subject
    m.set_content("\n\n".join(b.text for b in blocks if b.visible))
    m.add_alternative(to_html(blocks, subject), subtype="html")
    return bytes(m)


def to_png(text: str, width: int = 900) -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    lines = [ln for para in text.split("\n") for ln in (_wrap(para, 70) or [""])]
    try:
        font = ImageFont.truetype("arial.ttf", 22)
    except OSError:
        font = ImageFont.load_default()
    img = Image.new("RGB", (width, 40 + 32 * len(lines)), "white")
    draw = ImageDraw.Draw(img)
    for i, ln in enumerate(lines):
        draw.text((20, 20 + 32 * i), ln, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _wrap(text: str, width: int) -> list[str]:
    out: list[str] = []
    for para in text.split("\n"):
        line = ""
        for word in para.split():
            if len(line) + len(word) + 1 > width:
                out.append(line)
                line = word
            else:
                line = f"{line} {word}".strip()
        out.append(line)
    return out


RENDERERS = {"html": to_html, "md": to_markdown, "pdf": to_pdf, "docx": to_docx}
